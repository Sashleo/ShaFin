"""
reports/docx_renderer.py — Word rendering of a report model (reports/content.py).

Paragraphs, real Word tables and chart images; the same content as the PDF.
"""

import base64
import io

from docx import Document
from docx.shared import Cm, Pt


def render_docx(model: dict, generated_at: str) -> bytes:
    doc = Document()
    doc.core_properties.title = model["title"]
    doc.core_properties.author = "ShaFin"

    doc.add_heading(model["title"], level=0)
    doc.add_paragraph(f"{model['subtitle']} · generated {generated_at}")

    for section in model["sections"]:
        doc.add_heading(section["heading"], level=1)
        for block in section["blocks"]:
            kind = block[0]
            if kind == "p":
                doc.add_paragraph(block[1])
            elif kind == "table":
                headers, rows = block[1], block[2]
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, headers):
                    cell.text = str(text)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                for row in rows:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = str(text)
                doc.add_paragraph()
            elif kind == "chart":
                doc.add_picture(io.BytesIO(base64.b64decode(block[1])), width=Cm(16))
                caption = doc.add_paragraph(block[2])
                caption.runs[0].font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
